"""
合并报告生成器：结构化JSON → PDF 或 Word
灵活版：通过 sections 列表自由定义报告结构和顺序
"""
import os, uuid
from datetime import datetime


def _find_font():
    for p in [
        os.path.join(os.path.dirname(__file__), "..", "fonts", "SimHei.ttf"),
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\msyh.ttf",
    ]:
        if os.path.exists(p):
            return p
    return None


def _chart_image(chart_data: dict, output_dir: str) -> str:
    """生成图表图片，带输入校验和异常保护"""
    cats = chart_data.get("categories", [])
    vals = chart_data.get("values", [])
    if not cats or not vals or len(cats) != len(vals):
        raise ValueError(f"图表数据无效: categories={len(cats)}, values={len(vals)}")

    try:
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
    except ImportError:
        raise ImportError("matplotlib 未安装，请执行: pip install matplotlib")

    plt.rcParams['font.sans-serif'] = ['SimHei', 'Microsoft YaHei', 'Arial']
    plt.rcParams['axes.unicode_minus'] = False

    ct = chart_data.get("chart_type", "bar")
    title = chart_data.get("title", "")
    sn = chart_data.get("series_name", "")

    fp = os.path.join(output_dir, f"chart_{uuid.uuid4().hex[:8]}.png")
    fig, ax = plt.subplots(figsize=(8, 5))
    colors = ['#4472C4', '#ED7D31', '#A5A5A5', '#FFC000', '#5B9BD5', '#70AD47']

    if ct == "bar":
        bars = ax.bar(cats, vals, color=colors[:len(cats)], edgecolor='white', linewidth=0.5)
        ax.set_title(title, fontsize=14, fontweight='bold', pad=12)
        ax.set_ylabel(sn, fontsize=10)
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        max_val = max(vals) if vals else 1
        for b, v in zip(bars, vals):
            ax.text(b.get_x() + b.get_width() / 2., b.get_height() + max_val * 0.01,
                    str(v), ha='center', va='bottom', fontsize=9, fontweight='bold')
    elif ct == "pie":
        wedges, texts, autotexts = ax.pie(
            vals, labels=cats, autopct='%1.1f%%',
            colors=colors[:len(cats)], pctdistance=0.6,
            wedgeprops={'linewidth': 1, 'edgecolor': 'white'})
        for t in autotexts:
            t.set_fontsize(9)
            t.set_fontweight('bold')
        ax.set_title(title, fontsize=14, fontweight='bold', pad=12)
    elif ct == "line":
        ax.plot(range(len(cats)), vals, marker='o', color='#4472C4', linewidth=2.5, markersize=7)
        ax.fill_between(range(len(cats)), vals, alpha=0.1, color='#4472C4')
        ax.set_xticks(range(len(cats)))
        ax.set_xticklabels(cats, rotation=30, ha='right', fontsize=8)
        ax.set_title(title, fontsize=14, fontweight='bold', pad=12)
        ax.set_ylabel(sn, fontsize=10)
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        ax.grid(True, alpha=0.2, linestyle='--')
        for i, v in enumerate(vals):
            ax.annotate(str(v), (i, v), textcoords="offset points", xytext=(0, 12),
                       ha='center', fontsize=9, fontweight='bold')
    else:
        ax.barh(cats, vals, color=colors[:len(cats)], edgecolor='white')
        ax.set_title(title, fontsize=14, fontweight='bold', pad=12)
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)

    plt.tight_layout()
    plt.savefig(fp, dpi=150, bbox_inches='tight', facecolor='white')
    plt.close()
    return fp


def _auto_col_widths(pdf, headers, rows, pw):
    n = len(headers)
    w = [pdf.get_string_width(str(h)) + 6 for h in headers]
    for row in rows:
        for i in range(min(len(row), n)):
            cw = pdf.get_string_width(str(row[i])) + 6
            if cw > w[i]:
                w[i] = cw
    total = sum(w)
    if total == 0:
        return [pw / n] * n
    return [x / total * pw for x in w]


# ==================== PDF ====================

def _divider_line(pdf, pw, color=(200, 210, 220)):
    y = pdf.get_y()
    pdf.set_draw_color(*color)
    pdf.set_line_width(0.3)
    pdf.line(pdf.l_margin + pw * 0.05, y, pdf.l_margin + pw * 0.95, y)
    pdf.ln(3)


def _pdf_section_title(pdf, pw, text):
    pdf.ln(2)
    _divider_line(pdf, pw)
    pdf.set_font(pdf.font, 'B', 12)
    pdf.set_text_color(30, 50, 80)
    pdf.cell(pw, 7, text, 0, 1, 'L')
    pdf.set_text_color(0, 0, 0)
    pdf.ln(1)


def _clean_text(text: str) -> str:
    if not text:
        return text
    import re
    # 1. 只压缩水平空白字符（空格/制表符/不间断空格/Unicode排版空格等），保留换行符
    #    0x20=空格 0x09=制表符 \u00A0=不间断空格 \u2000-\u200A=Unicode排版空格（\u3000全角空格保留用于中文排版）
    text = re.sub(r'[ \t\u00A0\u2000-\u200A]+', ' ', text)
    # 2. 清理每行首尾空格
    text = re.sub(r'^ +', '', text, flags=re.MULTILINE)
    text = re.sub(r' +$', '', text, flags=re.MULTILINE)
    return text


def _pdf_render_text(pdf, pw, text):
    """Render text with proper multi-wrap handling for CJK mixed text."""
    pdf.set_font(pdf.font, '', 10)
    pdf.set_text_color(50, 55, 65)
    cleaned = _clean_text(text)
    for line in cleaned.split('\n'):
        if line.strip():
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(pw, 5.5, line, 0, 'L')
            # multi_cell with CJK font may drift x; set_x ensures proper alignment
        else:
            pdf.ln(5.5)
    pdf.set_text_color(0, 0, 0)


def _pdf_render_findings(pdf, pw, findings):
    for i, f in enumerate(findings):
        pdf.set_x(pdf.l_margin)
        pdf.set_font(pdf.font, 'B', 9)
        pdf.set_text_color(68, 114, 196)
        pdf.cell(8, 5, str(i + 1), 0, 0, 'R')
        pdf.set_font(pdf.font, '', 10)
        pdf.set_text_color(50, 55, 65)
        # 使用 multi_cell 正确处理多行文本换行，避免超长段落丢失
        indent_x = pdf.l_margin + 8
        pdf.set_x(indent_x)
        pdf.multi_cell(pw - 8, 5, _clean_text(f), 0, 'L')
        pdf.set_text_color(0, 0, 0)
    pdf.ln(2)


def _pdf_render_table(pdf, pw, headers, rows):
    cw = _auto_col_widths(pdf, headers, rows, pw)
    pdf.set_font(pdf.font, 'B', 8)
    pdf.set_fill_color(68, 114, 196)
    pdf.set_text_color(255, 255, 255)
    for i, hh in enumerate(headers):
        pdf.cell(cw[i], 7, str(hh), 1, 0, 'C', True)
    pdf.ln()
    pdf.set_draw_color(215, 220, 228)
    pdf.set_font(pdf.font, '', 8)
    pdf.set_text_color(50, 55, 65)
    for ri, rd in enumerate(rows):
        pdf.set_fill_color(245, 248, 252) if ri % 2 == 0 else pdf.set_fill_color(255, 255, 255)
        for i, cv in enumerate(rd):
            if i < len(cw):
                pdf.cell(cw[i], 6, str(cv)[:150], 1, 0, 'C', True)
        pdf.ln()
    pdf.set_text_color(0, 0, 0)
    pdf.ln(3)


def _pdf_render_chart(pdf, pw, cd, output_dir):
    cid = uuid.uuid4().hex[:8]
    cp = _chart_image(cd, output_dir)
    if os.path.exists(cp):
        w = min(170, pw)
        pdf.image(cp, x=(pdf.w - w) / 2, w=w)
        pdf.ln(2)
        if cd.get("title"):
            pdf.set_font(pdf.font, '', 8)
            pdf.set_text_color(140, 150, 160)
            pdf.cell(pw, 5, cd["title"], 0, 1, 'C')
            pdf.set_text_color(0, 0, 0)
    return cid


def _generate_pdf(summary: dict, output_dir: str, font_path: str) -> str:
    from fpdf import FPDF

    title = summary.get("title", "数据分析报告")
    subtitle = summary.get("subtitle", "")
    dept = summary.get("department", "")

    class PDF(FPDF):
        def __init__(self):
            super().__init__()
            if font_path:
                self.add_font("CN", "", font_path)
                self.add_font("CN", "B", font_path)
                self.font = "CN"
                self.set_char_spacing(0)
            else:
                self.font = "Helvetica"

        def header(self):
            if self.page_no() > 1:
                self.set_font(self.font, '', 7)
                self.set_text_color(160, 168, 178)
                self.cell(0, 6, title, 0, 1, 'R')
                self.set_text_color(0, 0, 0)
                self.ln(1)

        def footer(self):
            self.set_y(-15)
            self.set_font(self.font, '', 7)
            self.set_text_color(160, 168, 178)
            self.cell(0, 10, f'- {self.page_no()}/{{nb}} -', 0, 0, 'C')
            self.set_text_color(0, 0, 0)

    pdf = PDF()
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(True, 18)
    pdf.add_page()
    pw = pdf.w - pdf.l_margin - pdf.r_margin

    # ====== 封面 ======
    pdf.ln(4)
    pdf.set_draw_color(68, 114, 196)
    pdf.set_line_width(2)
    y0 = pdf.get_y()
    pdf.line(pdf.l_margin + pw * 0.25, y0, pdf.l_margin + pw * 0.75, y0)
    pdf.ln(6)

    pdf.set_font(pdf.font, 'B', 20)
    pdf.set_text_color(20, 40, 70)
    pdf.multi_cell(pw, 12, title, 0, 'C')
    pdf.set_x(pdf.l_margin)  # CJK multi_cell reset

    if subtitle:
        pdf.ln(2)
        pdf.set_x(pdf.l_margin)
        pdf.set_font(pdf.font, '', 11)
        pdf.set_text_color(100, 110, 130)
        pdf.multi_cell(pw, 7, subtitle, 0, 'C')

    pdf.ln(3)
    pdf.set_draw_color(68, 114, 196)
    pdf.set_line_width(0.8)
    y1 = pdf.get_y()
    pdf.line(pdf.l_margin + pw * 0.3, y1, pdf.l_margin + pw * 0.7, y1)
    pdf.ln(5)

    # 元信息
    ts = summary.get('timestamp', datetime.now().isoformat())[:19]
    src = summary.get('matched_count', 0)
    conf = summary.get('confidence', 0)
    parts = [f"Time: {ts}", f"Sources: {src}", f"Confidence: {conf:.0%}"]
    if dept:
        parts.insert(0, f"Dept: {dept}")
    pdf.set_font(pdf.font, '', 8)
    pdf.set_text_color(140, 150, 160)
    pdf.cell(pw, 5, "  |  ".join(parts), 0, 1, 'C')
    pdf.set_text_color(0, 0, 0)
    pdf.ln(4)

    # ====== 动态 sections ======
    sections = summary.get("sections", [])

    # 向后兼容：如果没有 sections，从旧字段构造
    if not sections:
        sections = []
        if summary.get("answer"):
            sections.append({"type": "text", "title": "内容", "content": summary["answer"]})
        if summary.get("key_findings"):
            sections.append({"type": "findings", "title": "关键发现", "items": summary["key_findings"]})
        if summary.get("table_data"):
            td = summary["table_data"]
            sections.append({"type": "table", "title": "数据总览", "headers": td.get("headers", []), "rows": td.get("rows", [])})
        if summary.get("chart_data"):
            sections.append({"type": "chart", "title": "数据图表", "data": summary["chart_data"]})
        if summary.get("sources"):
            sections.append({"type": "sources", "title": "数据来源", "items": summary["sources"]})

    chart_files_to_clean = []

    for sec in sections:
        stype = sec.get("type", "text")
        stitle = sec.get("title", "")

        if stype == "text":
            content = sec.get("content", "")
            if not content:
                continue
            if stitle:
                _pdf_section_title(pdf, pw, stitle)
            _pdf_render_text(pdf, pw, content)

        elif stype == "findings":
            items = sec.get("items", [])
            if not items:
                continue
            if stitle:
                _pdf_section_title(pdf, pw, stitle)
            _pdf_render_findings(pdf, pw, items)

        elif stype == "table":
            headers = sec.get("headers", [])
            rows = sec.get("rows", [])
            if not rows:
                continue
            if stitle:
                _pdf_section_title(pdf, pw, stitle)
            _pdf_render_table(pdf, pw, headers, rows)

        elif stype == "chart":
            cd = sec.get("data", {})
            if not cd.get("values"):
                continue
            if stitle:
                _pdf_section_title(pdf, pw, stitle)
            cid = _pdf_render_chart(pdf, pw, cd, output_dir)
            if cid:
                chart_files_to_clean.append(cid)

        elif stype == "sources":
            items = sec.get("items", [])
            if not items:
                continue
            if stitle:
                _pdf_section_title(pdf, pw, stitle)
            pdf.set_font(pdf.font, '', 9)
            pdf.set_text_color(120, 130, 140)
            for s in items:
                pdf.cell(pw, 5, f"  - {s}", 0, 1, 'L')
            pdf.set_text_color(0, 0, 0)

    # 页脚
    pdf.ln(4)
    pdf.set_draw_color(210, 218, 228)
    pdf.set_line_width(0.3)
    yf = pdf.get_y()
    pdf.line(pdf.l_margin + pw * 0.1, yf, pdf.l_margin + pw * 0.9, yf)
    pdf.ln(3)
    pdf.set_font(pdf.font, '', 7)
    pdf.set_text_color(170, 175, 185)
    pdf.cell(pw, 4, "Generated by 巨能环境文档安全管理系统 | Confidential", 0, 1, 'C')
    pdf.set_text_color(0, 0, 0)

    fn = f"report_{uuid.uuid4().hex[:8]}_{datetime.now().strftime('%Y%m%d')}.pdf"
    fp = os.path.join(output_dir, fn)
    pdf.output(fp)

    for cid in chart_files_to_clean:
        try: os.remove(os.path.join(output_dir, f"chart_{cid}.png"))
        except: pass
    return fp


# ==================== Word ====================

def _add_w_divider(doc):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), 'D2DAE4')
    pBdr.append(bottom)
    pPr.append(pBdr)


def _set_run_font(run, name, size=None, color=None, bold=False):
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    run.font.name = name
    run.element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    run.font.bold = bold


def _generate_word(summary: dict, output_dir: str, font_path: str) -> str:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)

    font_name = os.path.splitext(os.path.basename(font_path))[0] if font_path else 'SimHei'
    s = doc.styles['Normal']
    s.font.name = font_name
    s.font.size = Pt(10.5)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.line_spacing = 1.35
    s.element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

    title = summary.get("title", "数据分析报告")
    subtitle = summary.get("subtitle", "")
    dept = summary.get("department", "")

    # 封面
    doc.add_paragraph()
    t = doc.add_heading(title, 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in t.runs:
        _set_run_font(r, font_name, color=(20, 40, 70))

    if subtitle:
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = sp.add_run(subtitle)
        _set_run_font(sr, font_name, 12, color=(100, 110, 130))

    ts = summary.get('timestamp', datetime.now().isoformat())[:19]
    src = summary.get('matched_count', 0)
    conf = summary.get('confidence', 0)
    parts = [f"Time: {ts}", f"Sources: {src}", f"Confidence: {conf:.0%}"]
    if dept:
        parts.insert(0, f"Dept: {dept}")
    mp = doc.add_paragraph()
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = mp.add_run("  |  ".join(parts))
    _set_run_font(mr, font_name, 8, color=(140, 150, 160))

    _add_w_divider(doc)
    doc.add_paragraph()

    # ====== 动态 sections ======
    sections = summary.get("sections", [])

    # 向后兼容
    if not sections:
        sections = []
        if summary.get("answer"):
            sections.append({"type": "text", "title": "内容", "content": summary["answer"]})
        if summary.get("key_findings"):
            sections.append({"type": "findings", "title": "关键发现", "items": summary["key_findings"]})
        if summary.get("table_data"):
            td = summary["table_data"]
            sections.append({"type": "table", "title": "数据总览", "headers": td.get("headers", []), "rows": td.get("rows", [])})
        if summary.get("chart_data"):
            sections.append({"type": "chart", "title": "数据图表", "data": summary["chart_data"]})
        if summary.get("sources"):
            sections.append({"type": "sources", "title": "数据来源", "items": summary["sources"]})

    chart_files_to_clean = []

    for sec in sections:
        stype = sec.get("type", "text")
        stitle = sec.get("title", "")

        if stype == "text":
            content = sec.get("content", "")
            if not content:
                continue
            if stitle:
                h = doc.add_heading(stitle, 1)
                for r in h.runs:
                    _set_run_font(r, font_name, color=(30, 50, 80))
            p = doc.add_paragraph(_clean_text(content))
            for r in p.runs:
                _set_run_font(r, font_name, 10)

        elif stype == "findings":
            items = sec.get("items", [])
            if not items:
                continue
            if stitle:
                h = doc.add_heading(stitle, 1)
                for r in h.runs:
                    _set_run_font(r, font_name, color=(30, 50, 80))
            for i, f in enumerate(items):
                p = doc.add_paragraph()
                rn = p.add_run(f"{i + 1}. ")
                _set_run_font(rn, font_name, 10, color=(68, 114, 196), bold=True)
                rt = p.add_run(_clean_text(f))
                _set_run_font(rt, font_name, 10)

        elif stype == "table":
            headers = sec.get("headers", [])
            rows = sec.get("rows", [])
            if not rows:
                continue
            if stitle:
                h = doc.add_heading(stitle, 1)
                for r in h.runs:
                    _set_run_font(r, font_name, color=(30, 50, 80))
            ncols = len(headers)
            nrows = len(rows)
            tbl = doc.add_table(rows=1 + nrows, cols=ncols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.style = 'Light Grid Accent 1'
            tbl.autofit = False

            # 列宽
            page_width_cm = 16.0
            col_lens = [len(str(h)) for h in headers]
            for row in rows:
                for i, cell in enumerate(row):
                    if i < ncols:
                        cl = len(str(cell))
                        if cl > col_lens[i]:
                            col_lens[i] = cl
            total_len = sum(col_lens)
            if total_len == 0:
                col_widths = [page_width_cm / ncols] * ncols
            else:
                col_widths = [max(page_width_cm * (l / total_len), 1.5) for l in col_lens]
            total_w = sum(col_widths)
            if total_w > page_width_cm:
                scale = page_width_cm / total_w
                col_widths = [w * scale for w in col_widths]
            for i in range(ncols):
                tbl.columns[i].width = Cm(col_widths[i])

            for i, hh in enumerate(headers):
                c = tbl.rows[0].cells[i]
                c.text = str(hh)
                from docx.oxml import OxmlElement
                sh = c._element.get_or_add_tcPr()
                se = sh.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): '4472C4'})
                sh.append(se)
                for cp in c.paragraphs:
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for rn in cp.runs:
                        _set_run_font(rn, font_name, 9, color=(255, 255, 255), bold=True)

            for ri, rd in enumerate(rows):
                for ci, cv in enumerate(rd):
                    c = tbl.rows[ri + 1].cells[ci]
                    c.text = str(cv)[:200]
                    for cp in c.paragraphs:
                        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for rn in cp.runs:
                            _set_run_font(rn, font_name, 9)
                    if ri % 2 == 0:
                        from docx.oxml import OxmlElement
                        sh = c._element.get_or_add_tcPr()
                        se = sh.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): 'F5F8FC'})
                        sh.append(se)

        elif stype == "chart":
            cd = sec.get("data", {})
            if not cd.get("values"):
                continue
            if stitle:
                h = doc.add_heading(stitle, 1)
                for r in h.runs:
                    _set_run_font(r, font_name, color=(30, 50, 80))
            cid = uuid.uuid4().hex[:8]
            cp = _chart_image(cd, output_dir)
            if os.path.exists(cp):
                doc.add_picture(cp, width=Inches(5.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if cd.get("title"):
                    cap = doc.add_paragraph(cd["title"])
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for rn in cap.runs:
                        _set_run_font(rn, font_name, 9, color=(140, 150, 160))
                chart_files_to_clean.append(cid)

        elif stype == "sources":
            items = sec.get("items", [])
            if not items:
                continue
            if stitle:
                h = doc.add_heading(stitle, 1)
                for r in h.runs:
                    _set_run_font(r, font_name, color=(30, 50, 80))
            for s in items:
                p = doc.add_paragraph(s, style='List Bullet')
                for r in p.runs:
                    _set_run_font(r, font_name, 9, color=(120, 130, 140))

    # 页脚
    _add_w_divider(doc)
    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Generated by 巨能环境文档安全管理系统 | Confidential")
    _set_run_font(fr, font_name, 7, color=(170, 175, 185))

    fn = f"report_{uuid.uuid4().hex[:8]}_{datetime.now().strftime('%Y%m%d')}.docx"
    fp_path = os.path.join(output_dir, fn)
    doc.save(fp_path)

    for cid in chart_files_to_clean:
        try: os.remove(os.path.join(output_dir, f"chart_{cid}.png"))
        except: pass
    return fp_path


# ==================== 入口 ====================

def generate_report(summary: dict = None, format: str = "pdf", output_dir: str = None,
                    doc_id: str = None, answer: str = None, title: str = None,
                    fmt: str = None, chart_data: str = None, table_data: str = None) -> str:
    """Unified report generation entry point.
    Accepts: generate_report(summary_dict) OR generate_report(doc_id=..., answer=..., title=..., fmt=...)
    """
    import json as _json
    if summary is None or doc_id is not None:
        cd = None; td = None
        if chart_data:
            try: cd = _json.loads(chart_data) if isinstance(chart_data, str) else chart_data
            except Exception: pass
        if table_data:
            try: td = _json.loads(table_data) if isinstance(table_data, str) else table_data
            except Exception: pass
        summary = {
            "title": title or "数据分析报告",
            "subtitle": f"文档 {doc_id or ''} — AI智能总结",
            "department": "",
            "answer": answer or "",
            "matched_count": 1, "confidence": 1.0,
            "timestamp": datetime.now().isoformat(),
        }
        if cd: summary["chart_data"] = cd
        if td: summary["table_data"] = td
        if fmt and fmt != "pdf": format = fmt
    out = output_dir or os.path.join(os.getcwd(), "reports")
    os.makedirs(out, exist_ok=True)
    fp = _find_font()
    if not fp:
        raise RuntimeError(
            "未找到中文字体（SimHei/msyh）。报告生成已中止以避免乱码。请执行以下任一操作：\n"
            "1. 将 SimHei.ttf 放入 report-generator/fonts/ 目录\n"
            "2. 安装中文字体包到 C:\\Windows\\Fonts\\\n"
            "No Chinese font found (SimHei/msyh). Report generation aborted to prevent garbled output."
        )
    try:
        return _generate_word(summary, out, fp) if format == "word" else _generate_pdf(summary, out, fp)
    except ImportError as e:
        raise ImportError(f"Missing dependency: {e}. PDF needs fpdf2, Word needs python-docx. "
                         f"Install with: pip install fpdf2 python-docx")
